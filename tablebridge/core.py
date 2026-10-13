"""Explicit input contracts, deterministic calculation and template export."""
from __future__ import annotations
import csv, io, json, re, unicodedata, hashlib
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook

SCHEMAS = {
 'leave': ['record_id','employee_id','name','department','start_date','end_date','days','leave_type','status'],
 'balance': ['employee_id','name','balance'],
 'overtime': ['record_id','employee_id','name','department','date','hours','status'],
}
LABELS = {'record_id':'单据编号','employee_id':'工号','name':'姓名','department':'部门','start_date':'开始日期','end_date':'结束日期','days':'天数','leave_type':'假期类型','status':'审批状态','balance':'年假余额','date':'日期','hours':'小时数','source_id':'来源单据'}
ALIASES = {
 'record_id':['单据编号','单号','申请编号','请假单号','加班单号','record_id'],
 'employee_id':['工号','员工编号','人员编码','职员代码','职员编号','employee_id','emp_id'],
 'name':['姓名','员工姓名','人员姓名','name','employee_name'],
 'department':['部门','所属部门','组织名称','department','dept'],
 'start_date':['开始日期','起始日期','请假开始','start_date'],
 'end_date':['结束日期','截止日期','请假结束','end_date'],
 'days':['天数','请假天数','休假天数','请假时长（天）','days'],
 'leave_type':['假期类型','请假类别','休假类型','leave_type'],
 'status':['审批状态','状态','审核结果','status'],
 'balance':['年假余额','剩余年假','剩余天数','可用年假（天）','balance'],
 'date':['日期','加班日期','工作日期','date'],
 'hours':['小时数','加班时长','加班小时','加班时长（小时）','hours'],
 'source_id':['来源单据','来源申请编号','原单号','source_id'],
}
OUTPUTS = {'leave':['employee_id','name','department','start_date','end_date','leave_type','days','source_id'], 'overtime':['employee_id','name','department','date','hours','source_id']}

class ContractError(ValueError): pass

def clean(v):
 if v is None: return ''
 if isinstance(v, datetime): return v.isoformat(sep=' ')
 if isinstance(v, date): return v.isoformat()
 return unicodedata.normalize('NFKC',str(v)).replace('\ufeff','').replace('\u200b','').strip()

def safe_text(v):
 # CSV and XLSX exports must not interpret untrusted strings as formulas.
 s=str(v)
 return "'"+s if s.lstrip().startswith(('=','+','-','@','\t','\r')) else s

def number(v, positive=False):
 s=clean(v)
 if not re.fullmatch(r'\d+(?:\.\d{1,2})?',s): raise ContractError('数值须为非负十进制，最多两位小数，不接受单位、公式或科学计数法')
 try: n=Decimal(s)
 except InvalidOperation: raise ContractError('数值无法解析')
 if n>Decimal('1000000000'): raise ContractError('数值超过本原型精确导出范围 1000000000')
 if positive and n<=0: raise ContractError('时长必须大于零')
 return n

def day(v):
 if isinstance(v, datetime): return v.date()
 if isinstance(v, date): return v
 s=clean(v)
 # ISO timestamps with a midnight time and unambiguous Y/M/D only.
 if re.fullmatch(r'\d{4}-\d{2}-\d{2} 00:00:00',s): s=s[:10]
 s=s.replace('年','-').replace('月','-').replace('日','').replace('/','-').replace('.','-')
 if not re.fullmatch(r'\d{4}-\d{1,2}-\d{1,2}',s): raise ContractError('日期须为明确的年月日，例如 2026-09-01')
 try: return date(*map(int,s.split('-')))
 except ValueError: raise ContractError('日期不存在')

def identifier(v):
 # A number may already have lost significant zeroes in the source workbook.
 if not isinstance(v,str): raise ContractError('工号必须为文本；数字工号可能已丢失前导零，请核对源表')
 s=clean(v)
 if not s: raise ContractError('缺少工号')
 return s

def xlsx_load(raw):
 if len(raw)>10*1024*1024: raise ContractError('单文件上限 10 MB')
 try:
  with ZipFile(io.BytesIO(raw)) as z:
   if sum(i.file_size for i in z.infolist())>50*1024*1024: raise ContractError('工作簿解压体积超过 50 MB')
   if any('vbaProject' in n or '/externalLinks/' in n for n in z.namelist()): raise ContractError('不接受宏或外部链接')
  wb=load_workbook(io.BytesIO(raw),data_only=False)
 except ContractError: raise
 except Exception as e: raise ContractError('无效 XLSX 工作簿') from e
 if len(wb.worksheets)!=1: raise ContractError('本版本每个文件只接受一张工作表；请拆分或选定工作表后另存')
 ws=wb.active
 if ws.max_row>5020 or ws.max_column>60: raise ContractError('最多 5000 条数据、60 列')
 if ws.merged_cells.ranges: raise ContractError('请先展开合并单元格并使用单行表头')
 return wb

def read_table(raw, filename, header_row=1, template=False):
 if not isinstance(header_row,int) or not 1<=header_row<=20: raise ContractError('表头行须为 1 至 20')
 suffix=filename.lower().rsplit('.',1)[-1]
 if suffix=='xlsx':
  wb=xlsx_load(raw); ws=wb.active
  matrix=[]; flags=[]
  for row in ws.iter_rows():
   matrix.append([c.value for c in row]); flags.append(any(c.data_type=='f' or c.data_type=='e' for c in row))
 elif suffix in ('csv','tsv') and not template:
  if len(raw)>10*1024*1024: raise ContractError('单文件上限 10 MB')
  text=None
  for enc in ('utf-8-sig','gb18030'):
   try: text=raw.decode(enc); break
   except UnicodeDecodeError: pass
  if text is None: raise ContractError('CSV 编码须为 UTF-8 或 GB18030')
  matrix=list(csv.reader(io.StringIO(text),delimiter='\t' if suffix=='tsv' else ',')); flags=[False]*len(matrix)
  if len(matrix)>5020 or any(len(r)>60 for r in matrix): raise ContractError('最多 5000 条数据、60 列')
 else: raise ContractError('源文件支持 XLSX、CSV、TSV；目标模板仅支持 XLSX')
 while len(matrix)>header_row and not any(clean(v) for v in matrix[-1]): matrix.pop(); flags.pop()
 if len(matrix)<header_row: raise ContractError('指定表头行不存在')
 if len(matrix)-header_row>5000: raise ContractError('最多 5000 条数据')
 if flags[header_row-1]: raise ContractError('表头行含公式或 Excel 错误，请改为明确文本')
 headers=[clean(x) for x in matrix[header_row-1]]
 while headers and not headers[-1]: headers.pop()
 if not headers or any(not h for h in headers) or len(headers)!=len(set(headers)): raise ContractError('表头为空或重复，请先修正')
 if any(len(h)>200 for h in headers): raise ContractError('单个表头不得超过 200 字符')
 rows=[]
 for i,r in enumerate(matrix[header_row:],header_row+1):
  if len(r)>len(headers) and any(clean(v) for v in r[len(headers):]): raise ContractError(f'第 {i} 行有无表头的数据列')
  rows.append({'row':i,'values':dict(zip(headers,(r+[None]*len(headers))[:len(headers)])), 'formula':flags[i-1], 'blank':not any(clean(v) for v in r)})
 return {'filename':filename,'sheet':wb.active.title if suffix=='xlsx' else 'CSV','headers':headers,'rows':rows,'sha256':hashlib.sha256(raw).hexdigest(),'header_row':header_row}

def suggest(headers, fields):
 result={}; ambiguous=[]
 for field in fields:
  options=[h for h in headers if clean(h).lower() in [clean(a).lower() for a in ALIASES.get(field,[field])]]
  if len(options)==1: result[field]=options[0]
  elif len(options)>1: ambiguous.append(field)
 return result, ambiguous

def validate_mapping(table, mapping, role):
 if not isinstance(mapping,dict): raise ContractError('字段映射须为对象')
 fields=SCHEMAS[role]
 for f in fields:
  if not isinstance(mapping.get(f),str) or mapping[f] not in table['headers']: raise ContractError(f'{role} 缺少字段映射 {LABELS[f]}')
 if len(set(mapping[f] for f in fields))!=len(fields): raise ContractError('一个源列不能同时对应两个必需字段')

def validate_plan(plan):
 if not isinstance(plan,dict): raise ContractError('处理计划须为对象')
 allowed={'operation','department','start','end','eligible_types','approved_statuses','excluded_statuses','annual_code','personal_code','precision','assumption_confirmed'}
 if set(plan)-allowed: raise ContractError('计划包含不支持的规则；请使用允许的配置项')
 if plan.get('operation') not in OUTPUTS: raise ContractError('仅支持 leave 或 overtime 任务')
 p=dict(plan)
 for f in ('start','end'):
  p[f]=day(p[f]).isoformat() if p.get(f) else ''
 if bool(p['start'])!=bool(p['end']): raise ContractError('筛选起止日期须同时填写')
 if p['start'] and p['start']>p['end']: raise ContractError('筛选开始日期晚于结束日期')
 p['department']=clean(p.get('department',''))
 for f,default in [('eligible_types',['年假','事假']),('approved_statuses',['已通过','审批通过','已批准','approved']),('excluded_statuses',['待审批','已驳回','已撤销','pending','rejected'])]:
  v=p.get(f,default)
  if not isinstance(v,list) or not v or any(not isinstance(x,str) or not clean(x) for x in v): raise ContractError(f'{f} 必须为非空文本数组')
  p[f]=list(dict.fromkeys(clean(x) for x in v))
 if set(p['approved_statuses']) & set(p['excluded_statuses']): raise ContractError('通过与排除状态不得重叠')
 p['annual_code']=clean(p.get('annual_code','年假')); p['personal_code']=clean(p.get('personal_code','事假'))
 if p['operation']=='leave' and (not p['annual_code'] or not p['personal_code'] or p['annual_code']==p['personal_code']): raise ContractError('年假和事假导入编码须非空且不同')
 if p.get('precision',2)!=2: raise ContractError('本版本只接受 0.01 天或小时的精度，不静默舍入')
 p['precision']=2
 if not p.get('assumption_confirmed'): raise ContractError('请先确认业务规则和余额口径')
 return p

def execute(tables, mappings, plan):
 p=validate_plan(plan); mode=p['operation']
 roles=['leave','balance'] if mode=='leave' else ['overtime']
 for role in roles:
  if role not in tables: raise ContractError(f'缺少 {role} 源文件')
  validate_mapping(tables[role],mappings.get(role,{}),role)
 errors=[]; events=[]; ledger=[]; outputs=[]; blocked=set(); balances={}; counts={}
 def log(role,row,code,reason,kind='error',emp=''):
  item={'source':tables[role]['filename'],'sheet':tables[role]['sheet'],'row':row['row'],'employee_id':emp,'code':code,'reason':reason,'values':{k:clean(v) for k,v in row['values'].items()}}
  (errors if kind=='error' else events).append(item)
  counts[role][kind]+=1
  if kind=='error' and emp: blocked.add(emp)
 def records(role):
  counts[role]={'total':len(tables[role]['rows']),'blank':0,'error':0,'skip':0,'dedup':0,'valid':0}
  found=[]
  for row in tables[role]['rows']:
   if row['blank']: counts[role]['blank']+=1; continue
   d={f:row['values'].get(mappings[role][f]) for f in SCHEMAS[role]}; emp=clean(d.get('employee_id'))
   try:
    if row['formula']: raise ContractError('本行含公式或 Excel 错误；请核实后另存为值')
    d['employee_id']=identifier(d['employee_id']); emp=d['employee_id']
    for f in ('name',):
     d[f]=clean(d[f])
     if not d[f]: raise ContractError(f'缺少 {LABELS[f]}')
    if role=='balance': d['balance']=number(d['balance'])
    else:
     for f in ('record_id','department','status'):
      d[f]=clean(d[f])
      if not d[f]: raise ContractError(f'缺少 {LABELS[f]}')
     if role=='leave':
      d['start_date']=day(d['start_date']); d['end_date']=day(d['end_date']); d['days']=number(d['days'],True); d['leave_type']=clean(d['leave_type'])
      if d['start_date']>d['end_date']: raise ContractError('请假开始晚于结束')
      if not d['leave_type']: raise ContractError('缺少假期类型')
     else:
      d['date']=day(d['date']); d['hours']=number(d['hours'],True)
      if d['hours']>24: raise ContractError('单日加班超过 24 小时')
    found.append((d,row))
   except ContractError as e: log(role,row,'INVALID_ROW',str(e),emp=emp)
  # Group first, then isolate ALL conflicting copies rather than pick a winner.
  groups={}
  for d,row in found: groups.setdefault(d['employee_id'] if role=='balance' else d['record_id'],[]).append((d,row))
  unique=[]
  for key,group in groups.items():
   group.sort(key=lambda pair:pair[1]['row'])
   first=group[0][0]
   if any(d!=first for d,row in group[1:]):
    for d,row in group: log(role,row,'DUPLICATE_CONFLICT','同一工号余额或同一单据编号存在不同内容',emp=d['employee_id'])
   else:
    unique.append(group[0])
    for d,row in group[1:]: log(role,row,'DUPLICATE_IDENTICAL','重复记录已去重','dedup',d['employee_id'])
  return unique
 if mode=='leave':
  for d,row in records('balance'): balances[d['employee_id']]={**d,'source_row':row['row']}
 source=records(mode)
 # Identity conflicts within the task cannot be resolved by a name-based join.
 names={}
 for d,row in source: names.setdefault(d['employee_id'],set()).add(d['name'])
 blocked.update(e for e,n in names.items() if len(n)>1)
 candidates=[]
 for d,row in source:
  emp=d['employee_id']
  if d['status'] in p['excluded_statuses']: log(mode,row,'NOT_APPROVED','明确未通过的记录已排除','skip',emp); continue
  if d['status'] not in p['approved_statuses']: log(mode,row,'UNKNOWN_STATUS','未知审批状态，请确认后重算',emp=emp); continue
  if p['department'] and d['department']!=p['department']: log(mode,row,'FILTER_DEPARTMENT','不属于筛选部门','skip',emp); continue
  if mode=='leave' and d['leave_type'] not in p['eligible_types']: log(mode,row,'FILTER_TYPE','不属于已确认的结算假期类型','skip',emp); continue
  a,b=(d['start_date'],d['end_date']) if mode=='leave' else (d['date'],d['date'])
  if p['start']:
   start,end=day(p['start']),day(p['end'])
   if b<start or a>end: log(mode,row,'FILTER_DATE','不在筛选日期范围内','skip',emp); continue
   if a<start or b>end: log(mode,row,'PARTIAL_PERIOD','跨越筛选边界，无法按日期推算工作日天数，请拆单',emp=emp); continue
  if mode=='leave':
   if emp not in balances: log(mode,row,'MISSING_BALANCE','未找到该工号的期初年假余额',emp=emp); continue
   if d['name']!=balances[emp]['name']: log(mode,row,'IDENTITY_CONFLICT','工号对应姓名与余额表不一致',emp=emp); continue
  candidates.append((d,row))
 # No calculation happens until all employee-affecting errors are known.
 candidates.sort(key=lambda v:(v[0]['employee_id'], v[0].get('start_date',v[0].get('date')),v[0]['record_id']))
 remaining={e:d['balance'] for e,d in balances.items()}
 for d,row in candidates:
  emp=d['employee_id']
  if emp in blocked: log(mode,row,'EMPLOYEE_BLOCKED','该工号存在异常，隔离其全部候选记录以免余额结算失真',emp=emp); continue
  if mode=='overtime':
   outputs.append({**{f:clean(d[f]) for f in ('employee_id','name','department')},'date':d['date'].isoformat(),'hours':str(d['hours']),'source_id':d['record_id']})
   ledger.append({'employee_id':emp,'record_id':d['record_id'],'source_row':row['row'],'hours':str(d['hours'])})
  else:
   before=remaining[emp]; annual=min(before,d['days']); personal=d['days']-annual; after=before-annual; remaining[emp]=after
   assert annual+personal==d['days'] and before-annual==after and after>=0
   base={f:clean(d[f]) for f in ('employee_id','name','department')}; base.update(start_date=d['start_date'].isoformat(),end_date=d['end_date'].isoformat(),source_id=d['record_id'])
   for code,value in ((p['annual_code'],annual),(p['personal_code'],personal)):
    if value: outputs.append({**base,'leave_type':code,'days':str(value)})
   ledger.append({'employee_id':emp,'record_id':d['record_id'],'source_row':row['row'],'balance_source_row':balances[emp]['source_row'],'days':str(d['days']),'before':str(before),'annual':str(annual),'personal':str(personal),'after':str(after)})
  counts[mode]['valid']+=1
 if mode=='leave':
  for emp,d in balances.items():
   if emp in blocked:
    row=next(r for r in tables['balance']['rows'] if r['row']==d['source_row']); log('balance',row,'EMPLOYEE_BLOCKED','关联请假数据存在异常，余额记录仅保留复核',emp=emp)
   else: counts['balance']['valid']+=1
 totals={}
 for f in (['days','annual','personal'] if mode=='leave' else ['hours']): totals[f]=str(sum((Decimal(x[f]) for x in ledger),Decimal(0)))
 for role,c in counts.items(): assert c['total']==sum(c[k] for k in ('blank','error','skip','dedup','valid')),(role,c)
 return {'plan':p,'outputs':outputs,'ledger':ledger,'errors':errors,'events':events,'counts':counts,'totals':totals,'balances':{e:str(v) for e,v in remaining.items() if e not in blocked},'inputs':{r:{k:t[k] for k in ('filename','sheet','sha256','header_row')} for r,t in tables.items()},'mappings':mappings}

def export_template(raw, table, target_mapping, result):
 """Template columns are the only columns in the import workbook."""
 mode=result['plan']['operation']; allowed=OUTPUTS[mode]
 if not isinstance(target_mapping,dict) or set(target_mapping)!=set(table['headers']): raise ContractError('请为目标模板每一列明确映射字段或常量')
 mapped=[]
 for header,rule in target_mapping.items():
  if not isinstance(rule,dict) or set(rule) not in ({'field'},{'constant'}): raise ContractError('目标映射仅支持 field 或 constant')
  if 'field' in rule:
   if rule['field'] not in allowed: raise ContractError(f'不支持的目标字段 {rule["field"]}')
   mapped.append(rule['field'])
  elif not isinstance(rule['constant'],str): raise ContractError('固定值必须为文本')
 required=['employee_id','leave_type','days','start_date','end_date'] if mode=='leave' else ['employee_id','date','hours']
 if not set(required)<=set(mapped): raise ContractError('目标模板缺少工号、类型、日期或时长等必需输出列')
 wb=xlsx_load(raw); ws=wb.active; h=table['header_row']; cols=len(table['headers'])
 # Template data must be blank. Refuse to silently erase prior imports/formulas.
 if any(not r['blank'] for r in table['rows']): raise ContractError('目标模板数据区域须为空；请勿上传已有业务数据的工作簿')
 styles=[copy(ws.cell(h+1,c)._style) for c in range(1,cols+1)]
 for i,record in enumerate(result['outputs'],h+1):
  for c,header in enumerate(table['headers'],1):
   rule=target_mapping[header]; field=rule.get('field'); value=record[field] if field else rule['constant']; cell=ws.cell(i,c); cell._style=copy(styles[c-1])
   if field in ('days','hours'):
    cell.value=float(Decimal(value)); cell.number_format='0.00'
   elif field in ('date','start_date','end_date'):
    cell.value=day(value); cell.number_format=ws.cell(h+1,c).number_format if ws.cell(h+1,c).number_format not in ('General','@') else 'yyyy-mm-dd'
   else:
    cell.value=str(value); cell.data_type='s'; cell.number_format='@'
 # Retain template header, column widths, validation and sheet name.
 buffer=io.BytesIO(); wb.save(buffer); data=buffer.getvalue()
 check=load_workbook(io.BytesIO(data),data_only=False)
 if [clean(check.active.cell(h,c).value) for c in range(1,cols+1)]!=table['headers']: raise ContractError('导出表头校验失败')
 if any(c.data_type=='f' for row in check.active for c in row): raise ContractError('导出不得包含公式')
 return data

def csv_bytes(records, fields):
 s=io.StringIO(newline=''); w=csv.DictWriter(s,fields,extrasaction='ignore'); w.writeheader()
 for r in records: w.writerow({k:safe_text(json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v) for k,v in r.items()})
 return s.getvalue().encode('utf-8-sig')
