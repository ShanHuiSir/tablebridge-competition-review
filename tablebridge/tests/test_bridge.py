import sys, io, json, unittest, tempfile, threading, urllib.request, urllib.error, base64
from pathlib import Path
from copy import deepcopy
from decimal import Decimal
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from core import *
from planner import local_plan, create_plan
from unittest.mock import patch
from openpyxl import Workbook
import app

SAMPLES=Path(__file__).resolve().parents[1]/'samples'
def fixture(mode='leave'):
 files={'leave':'OA_请假_模拟.csv','balance':'HR_年假余额_模拟.csv'} if mode=='leave' else {'overtime':'ERP_加班_模拟.csv'}
 tables={r:read_table((SAMPLES/fn).read_bytes(),fn) for r,fn in files.items()}
 mappings={r:suggest(t['headers'],SCHEMAS[r])[0] for r,t in tables.items()}
 plan={'operation':mode,'start':'2026-09-01','end':'2026-09-30','assumption_confirmed':True}
 if mode=='overtime':plan['department']='研发部'
 return tables,mappings,plan
def run_fixture(mode='leave'):return execute(*fixture(mode))
def template(mode='leave'):
 fn='HR_'+('请假' if mode=='leave' else '加班')+'导入模板_模拟.xlsx';raw=(SAMPLES/fn).read_bytes();t=read_table(raw,fn,template=True);fm,_=suggest(t['headers'],OUTPUTS[mode]);return raw,t,{h:{'field':f} for f,h in fm.items()}

class BridgeTests(unittest.TestCase):
 def test_main_annual_personal(self):
  r=run_fixture();self.assertEqual(r['totals'],{'days':'6.05','annual':'3.30','personal':'2.75'});self.assertEqual(len(r['outputs']),6);self.assertEqual(r['balances']['0003'],'0.00')
 def test_independent_excel_formula_oracle(self):
  wb=load_workbook(SAMPLES/'独立公式核算_模拟.xlsx',data_only=True);rows=list(wb.active.values)[1:6];ledger=run_fixture()['ledger']
  for row,l in zip(rows,ledger):
   self.assertEqual(row[0],'工号 '+l['employee_id']);self.assertEqual(row[1],l['record_id'])
   for index,field in [(2,'days'),(3,'before'),(4,'annual'),(5,'personal'),(6,'after')]: self.assertEqual(Decimal(str(row[index])),Decimal(l[field]))
 def test_row_accounting(self):
  r=run_fixture()
  for c in r['counts'].values():self.assertEqual(c['total'],sum(c[k] for k in ('blank','error','skip','dedup','valid')))
  self.assertEqual(r['counts']['leave'],{'total':15,'blank':1,'error':6,'skip':2,'dedup':1,'valid':5})
 def test_bad_employee_does_not_consume_balance(self):
  r=run_fixture();self.assertNotIn('0004',r['balances']);self.assertFalse(any(x['employee_id']=='0004' for x in r['outputs']))
 def test_missing_balance_never_zero_fallback(self):
  r=run_fixture();self.assertTrue(any(e['code']=='MISSING_BALANCE' for e in r['errors']));self.assertFalse(any(x['employee_id']=='0005' for x in r['outputs']))
 def test_duplicate_conflicts_all_isolated(self):
  r=run_fixture();self.assertEqual(sum(e['code']=='DUPLICATE_CONFLICT' for e in r['errors']),2);self.assertFalse(any(x['source_id']=='L012' for x in r['outputs']))
 def test_perturbed_file_changes_output(self):
  t,m,p=fixture();t['balance']['rows'][0]['values'][m['balance']['balance']]='4.00';r=execute(t,m,p);self.assertEqual(r['totals']['annual'],'4.30');self.assertEqual(r['totals']['personal'],'1.75')
 def test_order_independent_input(self):
  t,m,p=fixture();t['leave']['rows'].reverse();r=execute(t,m,p);self.assertEqual(r['ledger'],run_fixture()['ledger'])
 def test_decimal_not_binary(self):
  self.assertEqual(number('0.10')+number('0.20'),Decimal('0.30'))
 def test_number_noise(self):
  self.assertEqual(number(' \u200b１.２５\ufeff '),Decimal('1.25'))
  for value in ['2天','1,000','-1','nan','inf','1e2','0.001','=1+1']:
   with self.subTest(value=value),self.assertRaises(ContractError):number(value)
 def test_date_formats_and_invalid(self):
  self.assertEqual(day('2026年9月2日').isoformat(),'2026-09-02')
  for value in ['09/02/2026','2026-02-30','46200']:
   with self.subTest(value=value),self.assertRaises(ContractError):day(value)
 def test_numeric_id_blocked(self):
  t,m,p=fixture();t['leave']['rows'][0]['values'][m['leave']['employee_id']]=1;r=execute(t,m,p);self.assertTrue(any('前导零' in e['reason'] for e in r['errors']))
 def test_unknown_status_blocks_employee(self):
  t,m,p=fixture();t['leave']['rows'][0]['values'][m['leave']['status']]='已阅';r=execute(t,m,p);self.assertFalse(any(x['employee_id']=='0001' for x in r['outputs']))
 def test_identity_conflict(self):
  t,m,p=fixture();t['balance']['rows'][0]['values'][m['balance']['name']]='同名另人';r=execute(t,m,p);self.assertFalse(any(x['employee_id']=='0001' for x in r['outputs']))
 def test_partial_period_isolated(self):
  t,m,p=fixture();t['leave']['rows'][0]['values'][m['leave']['start_date']]='2026-08-31';r=execute(t,m,p);self.assertTrue(any(e['code']=='PARTIAL_PERIOD' for e in r['errors']));self.assertFalse(any(x['employee_id']=='0001' for x in r['outputs']))
 def test_unknown_header_can_be_configured(self):
  t,m,p=fixture();old=m['balance']['balance'];t['balance']['headers'][2]='本期可抵扣数';
  for row in t['balance']['rows']:row['values']['本期可抵扣数']=row['values'].pop(old)
  m['balance']['balance']='本期可抵扣数';self.assertEqual(execute(t,m,p)['totals'],run_fixture()['totals'])
 def test_missing_mapping_rejected(self):
  t,m,p=fixture();m['leave'].pop('employee_id')
  with self.assertRaises(ContractError):execute(t,m,p)
 def test_unconfirmed_rules_rejected(self):
  t,m,p=fixture();p['assumption_confirmed']=False
  with self.assertRaises(ContractError):execute(t,m,p)
 def test_unapproved_and_outside_date_excluded(self):
  r=run_fixture();self.assertEqual({e['code'] for e in r['events']},{'DUPLICATE_IDENTICAL','NOT_APPROVED','FILTER_DATE'})
 def test_overtime_department_and_dates(self):
  r=run_fixture('overtime');self.assertEqual(r['totals']['hours'],'3.75');self.assertEqual(len(r['outputs']),2);self.assertEqual(len(r['errors']),1)
 def test_source_formula_isolation(self):
  t,m,p=fixture();t['leave']['rows'][0]['formula']=True; r=execute(t,m,p);self.assertFalse(any(x['employee_id']=='0001' for x in r['outputs']))
 def test_xlsx_template_text_ids_numbers_dates(self):
  raw,t,m=template();data=export_template(raw,t,m,run_fixture());ws=load_workbook(io.BytesIO(data)).active
  self.assertEqual(ws['A2'].value,'0001');self.assertEqual(ws['A2'].data_type,'s');self.assertEqual(ws['G2'].value,2.5);self.assertIsInstance(ws['D2'].value,datetime);self.assertEqual(ws.title,'请假导入');self.assertEqual(ws['A1'].value,'工号')
 def test_template_unknown_column_requires_constant(self):
  raw,t,m=template();m['工号']={'field':'unknown'}
  with self.assertRaises(ContractError):export_template(raw,t,m,run_fixture())
 def test_template_with_data_rejected(self):
  raw,t,m=template();wb=load_workbook(io.BytesIO(raw));wb.active['A2']='OLD';buf=io.BytesIO();wb.save(buf);raw=buf.getvalue();t=read_table(raw,'t.xlsx',template=True)
  with self.assertRaises(ContractError):export_template(raw,t,m,run_fixture())
 def test_csv_formula_protection(self):
  self.assertIn(b"'=1+1",csv_bytes([{'name':'=1+1'}],['name']))
 def test_table_input_contract(self):
  with self.assertRaises(ContractError):read_table('工号,工号\n0001,0002'.encode(),'x.csv')
  for header_row in (1,20):
   for count in (5000,5001):
    raw=('\n'*(header_row-1)+'工号\n'+'0001\n'*count).encode()
    wb=Workbook();ws=wb.active;ws.cell(header_row,1,'工号')
    for row in range(header_row+1,header_row+count+1):ws.cell(row,1,'0001')
    buf=io.BytesIO();wb.save(buf)
    for name,data in [('x.csv',raw),('x.xlsx',buf.getvalue())]:
     with self.subTest(header_row=header_row,count=count,name=name):
      if count==5000:self.assertEqual(len(read_table(data,name,header_row)['rows']),5000)
      else:
       with self.assertRaises(ContractError):read_table(data,name,header_row)
 def test_bad_plan_operators_rejected(self):
  t,m,p=fixture();p['python']='exec()'
  with self.assertRaises(ContractError):execute(t,m,p)
 def test_local_natural_language(self):
  r=local_plan('筛选加班。部门：研发部；2026-09-01 至 2026-09-30。');self.assertEqual(r['plan']['department'],'研发部');self.assertFalse(r['plan']['assumption_confirmed'])
 def test_ambiguous_natural_language(self):
  for prompt in ['处理一下表格','优先扣年假，按姓名关联','筛选本月加班','筛选加班，研发部门','筛选加班 2026-09-01']:
   with self.subTest(prompt=prompt),self.assertRaises(ContractError):local_plan(prompt)

class ApiTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory();app.DATA=Path(cls.tmp.name);cls.server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler);cls.base=f'http://127.0.0.1:{cls.server.server_port}';cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
 @classmethod
 def tearDownClass(cls):cls.server.shutdown();cls.server.server_close();cls.tmp.cleanup()
 def post(self,path,obj,origin=None):
  req=urllib.request.Request(self.base+'/api/'+path,data=json.dumps(obj).encode(),headers={'Content-Type':'application/json',**({'Origin':origin} if origin else {})});return json.load(urllib.request.urlopen(req))
 def test_roundtrip_download_and_immutable_history(self):
  sid=self.post('session',{})['session_id'];d=self.post('sample',{'session_id':sid});p=fixture()[2]
  body={'session_id':sid,'plan':p,'mappings':d['mappings'],'target_mapping':d['target_suggestions'],'confirmed':True};r=self.post('run',body)
  before=urllib.request.urlopen(self.base+r['downloads']['import.xlsx']).read();body['plan']['department']='研发部';r2=self.post('run',body)
  self.assertNotEqual(r['run_id'],r2['run_id']);self.assertEqual(before,urllib.request.urlopen(self.base+r['downloads']['import.xlsx']).read());self.assertEqual(r['totals']['annual'],'3.30')
 def test_cross_origin_rejected(self):
  with self.assertRaises(urllib.error.HTTPError):self.post('session',{},origin='https://example.com')
 def test_no_confirmation_no_export(self):
  sid=self.post('session',{})['session_id']
  with self.assertRaises(urllib.error.HTTPError):self.post('run',{'session_id':sid,'confirmed':False})

class PlannerTests(unittest.TestCase):
 def response(self,payload):
  class Response:
   def __enter__(self):return self
   def __exit__(self,*args):pass
   def read(self,n):return json.dumps(payload).encode()
  return Response()
 def test_model_request_contains_only_instruction_and_headers(self):
  t,m,p=fixture();p['assumption_confirmed']=True;payload={'choices':[{'message':{'content':json.dumps({'plan':p,'mappings':m,'warnings':[]})}}],'usage':{'total_tokens':123},'id':'mock-response','model':'mock-model'}
  with patch('planner.load_env'),patch.dict('os.environ',{'BRIDGE_API_KEY':'test-key','BRIDGE_MODEL':'mock-model','BRIDGE_BASE_URL':'https://example.test/v1'}),patch('urllib.request.urlopen',return_value=self.response(payload)) as request:
   result=create_plan('优先扣减年假，不足转事假',t,True);sent=json.loads(request.call_args.args[0].data)
   user=json.loads(sent['messages'][1]['content']);self.assertEqual(set(user),{'instruction','headers'});self.assertNotIn('陈曦',json.dumps(sent,ensure_ascii=False));self.assertNotIn('0001',json.dumps(sent));self.assertFalse(result['plan']['assumption_confirmed']);self.assertEqual(result['usage']['total_tokens'],123)
 def test_model_http_failure_no_local_fallback(self):
  t,m,p=fixture()
  with patch('planner.load_env'),patch.dict('os.environ',{'BRIDGE_API_KEY':'test-key'}),patch('urllib.request.urlopen',side_effect=urllib.error.HTTPError('https://example.test',401,'Unauthorized',{},None)):
   with self.assertRaisesRegex(ContractError,'HTTP 401'):create_plan('优先扣减年假，不足转事假',t,True)
 def test_model_invalid_json_no_execution(self):
  t,m,p=fixture()
  with patch('planner.load_env'),patch.dict('os.environ',{'BRIDGE_API_KEY':'test-key'}),patch('urllib.request.urlopen',return_value=self.response({'choices':[{'message':{'content':'not json'}}]})):
   with self.assertRaises(ContractError):create_plan('优先扣减年假，不足转事假',t,True)

if __name__=='__main__':unittest.main(verbosity=2)
